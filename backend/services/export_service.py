"""Gated DOCX export - refuses while any mandatory finding is open."""
from __future__ import annotations
import io
from typing import Any
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.core.rule_engine import all_rules
from backend.logging_setup import get_logger

log = get_logger("export")


class ExportBlocked(Exception): pass


def export_docx(result: dict[str, Any]) -> bytes:
    sc = result["scorecard"]
    if not sc.get("can_finalize"):
        log.warning("export blocked - %d mandatory findings", sc.get("mandatory_open", 0))
        raise ExportBlocked(f"{sc.get('mandatory_open',0)} mandatory compliance issue(s) still open.")
    brief = result["brief"]; dt = (brief.get("doc_type") or "RFP").upper()
    doc = Document()
    doc.add_heading(brief.get("title", dt), level=0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{dt}  \u00b7  Government of Uttarakhand").bold = True
    m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    money = result.get("derived", {}).get("money", {})
    m.add_run(f"Estimated cost: {money.get('project_cost','-')}   |   "
              f"Duration: {brief.get('duration_months','-')} months   |   "
              f"Compliance: {sc['overall_percent']}%")
    doc.add_paragraph()
    for val in result.get("sections", {}).values():
        doc.add_heading(val["title"], level=1)
        for line in str(val["body"]).split("\n"):
            line = line.strip()
            if not line: continue
            if line.startswith(("-", "*", "\u2022")):
                doc.add_paragraph(line.lstrip("-*\u2022 ").strip(), style="List Bullet")
            else:
                doc.add_paragraph(line)
    doc.add_page_break()
    doc.add_heading("Compliance Scorecard", level=1)
    t = doc.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
    h = t.rows[0].cells; h[0].text, h[1].text, h[2].text = "Framework", "Score", "Status"
    for row in sc["rows"]:
        c = t.add_row().cells
        c[0].text = row["framework"]
        c[1].text = f"{row['percent']}% ({row['passed']}/{row['total']})"
        c[2].text = row["status"].upper()
    doc.add_heading("Compliance Basis (Statutory References)", level=2)
    seen = set()
    for r in all_rules(doc_type=dt):
        k = (r.get("framework",""), r.get("citation",""))
        if k[1] and k not in seen:
            doc.add_paragraph(f"{k[0]}: {k[1]}", style="List Bullet"); seen.add(k)
    buf = io.BytesIO(); doc.save(buf)
    log.info("exported %s (%d sections)", dt, len(result.get("sections", {})))
    return buf.getvalue()
